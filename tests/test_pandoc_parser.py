"""Tests for pandoc conversion pipeline (epub/docx/html/tex → Markdown).

覆盖三块：
  1. cleaner.fix_hard_wraps / clean_text 的断行与实体清理；
  2. pandoc_parser：格式集合、TOC 提取、章节切块（含边界与兜底）；
  3. 引擎集成：epub 走「pandoc + TOC digest + 章节切块」路径。

pandoc / ebook-convert 是系统二进制，集成用例在缺失时自动跳过。
"""
from __future__ import annotations

import os

import pytest

from agent.rag.cleaner import clean_text, fix_hard_wraps
from agent.rag.document_outline import Chapter
from agent.rag import pandoc_parser as pp


# ---------------------------------------------------------------------------
# cleaner.fix_hard_wraps
# ---------------------------------------------------------------------------

def test_fix_hard_wraps_joins_mid_sentence_lines():
    text = "第一行没有结束\n继续第二行。\n第三段还没完\n第四行就完了。"
    out = fix_hard_wraps(text)
    assert "第一行没有结束 继续第二行。" in out
    assert "第三段还没完 第四行就完了。" in out
    # 两个段落之间保留换行
    assert "继续第二行。\n第三段还没完" in out


def test_fix_hard_wraps_preserves_blank_line_paragraphs():
    out = fix_hard_wraps("句子A未完\n句子B完。\n\n句子C完。")
    assert out == "句子A未完 句子B完。\n\n句子C完。"


def test_fix_hard_wraps_normalizes_crlf():
    out = fix_hard_wraps("第一行未完\r\n第二行完。\r\n\r\n第三段。")
    assert "\r" not in out
    assert "第一行未完 第二行完。" in out


def test_fix_hard_wraps_replaces_br_with_space():
    out = fix_hard_wraps("前半句<br>后半句\n另一段<br/>结束。")
    assert "前半句 后半句" in out
    assert "另一段 结束。" in out


def test_fix_hard_wraps_strips_known_entities():
    out = fix_hard_wraps("宽度&nbsp不足\n破折&mdash号完结。")
    assert "&nbsp" not in out
    assert "&mdash" not in out


def test_fix_hard_wraps_strips_gfm_hard_break_backslash():
    # pandoc gfm 把 <br> 输出为行尾反斜杠
    out = fix_hard_wraps("第一句完。\\\n第二句未完\\\n第三句完。")
    assert "\\" not in out
    assert "第二句未完 第三句完。" in out


def test_fix_hard_wraps_protects_markdown_structure():
    text = (
        "# 标题一\n"
        "段落未完\n段落完。\n"
        "- 列表项一\n"
        "- 列表项二\n"
        "下一行未完\n接上了。\n"
        "| 表头 | 列二 |\n"
        "| --- | --- |\n"
        "| a | b |\n"
        "```python\n"
        "# 注释 不 动\n"
        "x = 1\n"
        "```\n"
    )
    out = fix_hard_wraps(text)
    lines = out.split("\n")
    assert "# 标题一" in lines
    assert "- 列表项一" in lines
    assert "- 列表项二" in lines
    assert "| 表头 | 列二 |" in lines
    # 代码块逐字保留
    assert "# 注释 不 动" in out
    assert "x = 1" in out
    # 段落拼接正常
    assert "段落未完 段落完。" in out
    assert "下一行未完 接上了。" in out


def test_fix_hard_wraps_english_sentences():
    out = fix_hard_wraps("This line is not done\nyet. Next line\nis fine.")
    assert "This line is not done yet." in out
    assert "Next line\nis fine." not in out  # 被拼成一行
    assert "yet. Next line" in out or "yet.  Next line" in out


def test_clean_text_entity_without_semicolon():
    out = clean_text("A&nbsp词&nbsp; B&mdash词&mdash; C")
    assert "&nbsp" not in out
    assert "&mdash" not in out
    # 正文的 & 不受影响（非已知实体名）
    out2 = clean_text("R&D labs")
    assert "R&D" in out2


def test_clean_text_br_not_gluing_words():
    out = clean_text("<p>第一行</p><br><p>第二行</p>")
    assert "第一行 第二行" in out


def test_clean_text_nbsp():
    out = clean_text("词\u00a0汇")
    assert "\u00a0" not in out


# ---------------------------------------------------------------------------
# pandoc_parser：格式集合与可用性
# ---------------------------------------------------------------------------

def test_supported_extensions():
    assert pp.PANDOC_EBOOK_EXTS == {".epub", ".mobi", ".azw3", ".azw"}
    assert {".docx", ".html", ".htm", ".tex", ".latex"} <= pp.PANDOC_DOC_EXTS


def test_find_ebook_convert_macos_app_bundle():
    # 本机装了 calibre.app 时应能找到；没装时返回 None（不抛异常）
    result = pp.find_ebook_convert()
    assert result is None or os.path.isfile(result)


# ---------------------------------------------------------------------------
# 章节切块
# ---------------------------------------------------------------------------

_MD = """# 测试之书

前言内容。

# 第一章 起点

起点正文。 """ + "长" * 50 + """

## 第一节 细节

细节内容。

# 第二章 转折

转折正文。
"""


def test_build_chapter_chunks_with_toc_boundaries():
    toc = [Chapter(1, "第一章 起点", ""), Chapter(1, "第二章 转折", "")]
    chunks = pp.build_chapter_chunks(_MD, toc, max_chars=2000)
    # 前言+书名一组，第一章、第二章各一组（「# 测试之书」不在 TOC 顶级里，
    # 但它后面的前言非空 → 与书名合并为一组）
    assert len(chunks) == 3
    assert chunks[0].startswith("# 测试之书")
    assert "前言内容。" in chunks[0]
    assert chunks[1].startswith("# 第一章 起点")
    # 子标题行保留在章节文本里
    assert "## 第一节 细节" in chunks[1]
    assert chunks[2].startswith("# 第二章 转折")


def test_build_chapter_chunks_fallback_level_grouping():
    chunks = pp.build_chapter_chunks(_MD, [], max_chars=2000)
    # 无 TOC → 层级兜底：H1（书名/章）为界
    assert len(chunks) == 3
    assert chunks[0].startswith("# 测试之书")
    assert chunks[1].startswith("# 第一章 起点")
    assert chunks[2].startswith("# 第二章 转折")


def test_build_chapter_chunks_oversized_split():
    calls = []

    def fake_split(text):
        calls.append(text)
        return ["piece-1", "piece-2"]

    toc = [Chapter(1, "第一章 起点", "")]
    chunks = pp.build_chapter_chunks(_MD, toc, max_chars=50, oversized_split=fake_split)
    # 超长章节交给 oversized_split；短章节保持原样
    assert any("piece-1" in c or c == "piece-1" for c in chunks) or calls
    assert calls, "超长章节应触发兜底切片"
    # 未超长的章节不进入兜底
    assert any(c.startswith("# 测试之书") or "前言内容" in c for c in chunks)


def test_build_chapter_chunks_empty_input():
    assert pp.build_chapter_chunks("", None) == []


def test_split_sections_keeps_prefix_and_subheadings():
    sections = pp._split_sections(_MD)
    # 首行即 H1（无前缀段）：sections[0] = H1 书名，body 是前言
    assert sections[0][0] == 1 and "前言内容" in sections[0][2]
    levels = [s[0] for s in sections]
    assert 1 in levels and 2 in levels
    # 带前缀的情况：首个标题前有正文
    sections2 = pp._split_sections("开头的前言。\n\n# 第一章\n\n正文。")
    assert sections2[0] == (0, "", "开头的前言。")


# ---------------------------------------------------------------------------
# 真实 epub 集成（ebooklib 生成 → pandoc 转换）
# ---------------------------------------------------------------------------

ebooklib = pytest.importorskip("ebooklib")


def _make_epub(path: str, nested_toc: bool = False) -> None:
    from ebooklib import epub

    book = epub.EpubBook()
    book.set_identifier("test-book")
    book.set_title("测试之书")
    book.set_language("zh")
    book.add_author("作者")

    c0 = epub.EpubHtml(title="封面", file_name="cover.xhtml", lang="zh")
    c0.content = ("<html><body><h1>测试之书</h1>"
                  "<p>这是前言。一共两句话。<br>断行也要修好&nbsp好。</p></body></html>")
    c1 = epub.EpubHtml(title="第一章 起点", file_name="ch1.xhtml", lang="zh")
    c1.content = ("<html><body><h1>第一章 起点</h1>"
                  "<p>少年站在山门前，望着云海出神。</p></body></html>")
    c2 = epub.EpubHtml(title="第二章 转折", file_name="ch2.xhtml", lang="zh")
    c2.content = ("<html><body><h1>第二章 转折</h1>"
                  "<p>夜雨敲窗。</p><p>故事的走向变了&mdash没有人知道。</p></body></html>")
    for c in (c0, c1, c2):
        book.add_item(c)
    book.add_item(epub.EpubNcx())
    book.add_item(epub.EpubNav())
    if nested_toc:
        book.toc = (
            epub.Link("cover.xhtml", "封面", "cover"),
            (epub.Section("第一卷"), [
                epub.Link("ch1.xhtml", "第一章 起点", "ch1"),
                epub.Link("ch2.xhtml", "第二章 转折", "ch2"),
            ]),
        )
    else:
        book.toc = (
            epub.Link("cover.xhtml", "封面", "cover"),
            epub.Link("ch1.xhtml", "第一章 起点", "ch1"),
            epub.Link("ch2.xhtml", "第二章 转折", "ch2"),
        )
    book.spine = [c0, c1, c2]
    epub.write_epub(path, book)


def test_extract_epub_toc_flat_and_nested(tmp_path):
    _make_epub(str(tmp_path / "flat.epub"))
    toc = pp.extract_epub_toc(str(tmp_path / "flat.epub"))
    titles = [c.title for c in toc]
    assert "第一章 起点" in titles and "第二章 转折" in titles

    _make_epub(str(tmp_path / "nested.epub"), nested_toc=True)
    toc2 = pp.extract_epub_toc(str(tmp_path / "nested.epub"))
    by_title = {c.title: c.level for c in toc2}
    assert by_title["第一卷"] == 1
    assert by_title["第一章 起点"] == 2


@pytest.mark.skipif(not pp.is_pandoc_available(), reason="pandoc 未安装")
def test_convert_ebook_epub(tmp_path):
    _make_epub(str(tmp_path / "book.epub"))
    md, toc, chapters = pp.convert_ebook(str(tmp_path / "book.epub"))
    assert md and "第一章 起点" in md
    assert any(c.title == "第一章 起点" for c in toc)
    # 锚点 span 应在 convert_ebook 内被剥离；畸形实体由 engine 的
    # fix_hard_wraps 清理（fix_hard_wraps_strips_known_entities 覆盖）
    assert '<span id=' not in md


@pytest.mark.skipif(not pp.is_pandoc_available(), reason="pandoc 未安装")
def test_split_chapters_by_anchors_real_epub(tmp_path):
    """epub 直接读到准确章节：边界 = 内容文档锚点，标题 = TOC。"""
    _make_epub(str(tmp_path / "book.epub"))
    md, toc, chapters = pp.convert_ebook(str(tmp_path / "book.epub"))
    assert chapters is not None
    assert len(chapters) == 3  # 封面(书名+前言) / 第一章 / 第二章
    assert chapters[0].startswith("# 测试之书")
    assert "这是前言。" in chapters[0]
    assert "第一章 起点" in chapters[1] and "少年站在山门前" in chapters[1]
    assert "第二章 转折" in chapters[2] and "夜雨敲窗" in chapters[2]
    for ch in chapters:
        assert '<span id=' not in ch


def test_split_chapters_by_anchors_failure_fallbacks():
    md = '<span id="a.xhtml"></span>\n\n# 章\n\n正文。\n\n<span id="b.xhtml"></span>\n\n# 另一章\n\n内容。'
    # TOC 为空 → 无法提取准确目录 → 回退
    assert pp.split_chapters_by_anchors(md, []) is None
    # 无锚点 → 回退
    toc = [Chapter(1, "章", "a.xhtml")]
    assert pp.split_chapters_by_anchors("没有锚点的正文。", toc) is None
    # 单锚点（单文档书）→ 回退
    assert pp.split_chapters_by_anchors('<span id="a.xhtml"></span>\n\n# 章\n\n正文。', toc) is None


def test_split_chapters_by_anchors_toc_title_fills_missing_heading():
    md = (
        '<span id="a.xhtml"></span>\n\n# 第一章 起点\n\n甲。\n\n'
        '<span id="b.xhtml"></span>\n\n无标题文档的正文。\n\n'
    )
    toc = [
        Chapter(1, "第一章 起点", "a.xhtml"),
        Chapter(1, "第二章 来自目录的标题", "b.xhtml#sec1"),
    ]
    chunks = pp.split_chapters_by_anchors(md, toc)
    assert chunks is not None and len(chunks) == 2
    # b.xhtml 段无标题行 → 用 TOC 标题补上（href 带锚点 fragment 也能对齐）
    assert chunks[1].startswith("# 第二章 来自目录的标题")
    assert "无标题文档的正文" in chunks[1]


@pytest.mark.skipif(not pp.is_pandoc_available(), reason="pandoc 未安装")
def test_convert_document_html(tmp_path):
    f = tmp_path / "page.html"
    f.write_text("<html><body><h1>标题</h1><p>正文第一行<br>第二行</p></body></html>",
                 encoding="utf-8")
    md = pp.convert_document(str(f))
    assert md and "# 标题" in md and "正文第一行" in md


@pytest.mark.skipif(pp.find_ebook_convert() is None, reason="calibre ebook-convert 未安装")
def test_convert_ebook_mobi_via_calibre(tmp_path):
    epub_path = str(tmp_path / "src.epub")
    _make_epub(epub_path)
    mobi_path = str(tmp_path / "book.mobi")
    import subprocess
    res = subprocess.run(
        [pp.find_ebook_convert(), epub_path, mobi_path],
        capture_output=True, timeout=120,
    )
    if res.returncode != 0 or not os.path.exists(mobi_path):
        pytest.skip("ebook-convert 无法从测试 epub 生成 mobi")
    md, toc, chapters = pp.convert_ebook(mobi_path)
    assert md and len(md) > 50
    # TOC 从中转 epub 读出
    assert any(c.title == "第一章 起点" for c in toc)
    assert chapters is not None and len(chapters) >= 2


# ---------------------------------------------------------------------------
# extract_text 对 .tex 的路由
# ---------------------------------------------------------------------------

def test_supported_extensions_includes_tex():
    from agent.rag.parsers import SUPPORTED_EXTENSIONS
    assert ".tex" in SUPPORTED_EXTENSIONS and ".latex" in SUPPORTED_EXTENSIONS


def test_extract_text_tex_uses_pandoc(tmp_path, monkeypatch):
    f = tmp_path / "a.tex"
    f.write_text("\\documentclass{article}\\begin{document}hi\\end{document}")
    monkeypatch.setattr(pp, "convert_document", lambda p: "# converted")
    from agent.rag.parsers import extract_text
    assert extract_text(str(f)) == "# converted"


def test_extract_text_tex_missing_pandoc_message(tmp_path, monkeypatch):
    f = tmp_path / "a.tex"
    f.write_text("x")
    monkeypatch.setattr(pp, "convert_document", lambda p: None)
    monkeypatch.setattr(pp, "is_pandoc_available", lambda: False)
    from agent.rag.parsers import extract_text
    with pytest.raises(ValueError, match="brew install pandoc"):
        extract_text(str(f))


# ---------------------------------------------------------------------------
# 引擎集成：epub 走 TOC digest + 章节切块路径
# ---------------------------------------------------------------------------

@pytest.mark.skipif(not pp.is_pandoc_available(), reason="pandoc 未安装")
def test_ingest_epub_uses_toc_digest_and_chapter_chunks(tmp_path, fake_ef):
    from agent.rag.engine import RAGEngine

    kb = tmp_path / "kb"
    kb.mkdir()
    _make_epub(str(kb / "book.epub"), nested_toc=True)

    ws = tmp_path / "ws"
    eng = RAGEngine(workspace=str(ws), knowledge_base=str(kb),
                    embedding_function=fake_ef, chunk_size=500)
    stats = eng.ingest()

    assert stats["files_extracted"] == 1
    assert stats["errors"] == []
    # 封面/前言/两章 → 章节切块产出 3 片
    assert stats["chunks"] == 3

    # digest 来自 ebooklib TOC：「第一卷」是卷级 Section，只存在于目录树，
    # 不可能来自正文 Markdown 标题提取。
    outline = eng.get_document_outline("book.epub")
    assert outline is not None
    assert "第一卷" in outline["digest"]
    assert "第一章 起点" in outline["digest"]

    # markdown 缓存已写入（pandoc 输出经 fix_hard_wraps 清洗）
    docs = eng.list_documents()
    assert "book.epub" in docs


@pytest.mark.skipif(not pp.is_pandoc_available(), reason="pandoc 未安装")
def test_ingest_docx_via_pandoc(tmp_path, fake_ef):
    """docx 在 pandoc 可用时优先走 pandoc（is_markdown 路径）。"""
    from docx import Document
    from agent.rag.engine import RAGEngine

    kb = tmp_path / "kb"
    kb.mkdir()
    doc = Document()
    doc.add_heading("Pandoc 标题", level=1)
    doc.add_paragraph("这是正文内容。")
    doc.save(str(kb / "doc.docx"))

    ws = tmp_path / "ws"
    eng = RAGEngine(workspace=str(ws), knowledge_base=str(kb),
                    embedding_function=fake_ef)
    stats = eng.ingest()
    assert stats["files_extracted"] == 1
    assert stats["errors"] == []
    outline = eng.get_document_outline("doc.docx")
    assert outline is not None
    assert "Pandoc 标题" in outline["digest"]
